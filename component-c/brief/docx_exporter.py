"""
Component C - Function 2
Professional DOCX Legal Brief Exporter

Converts an existing LegalBrief into a styled
Microsoft Word document.

No Gemini calls or new legal claims are made.
"""

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from brief.output_models import LegalBrief
from brief.brief_formatter import format_legal_brief


# ---------------------------------------------
# Document theme
# ---------------------------------------------

NAVY = RGBColor(29, 49, 70)
TEAL = RGBColor(23, 104, 112)
SLATE = RGBColor(85, 99, 112)

DISCLAIMER = (
    "AI-assisted legal research draft. "
    "The cited authorities and legal conclusions "
    "must be independently verified before use."
)


# ---------------------------------------------
# Decorative paragraph border
# ---------------------------------------------

def _bottom_border(
    paragraph,
    color="176870",
    size="9",
):
    """
    Add a horizontal line below a paragraph.
    """

    p_pr = paragraph._p.get_or_add_pPr()

    borders = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")

    attributes = {
        "val": "single",
        "sz": size,
        "color": color,
        "space": "7",
    }

    for key, value in attributes.items():
        bottom.set(
            qn(f"w:{key}"),
            value,
        )

    borders.append(bottom)
    p_pr.append(borders)


# ---------------------------------------------
# Highlighted cell background
# ---------------------------------------------

def _shade_cell(
    cell,
    color="EDF3F5",
):
    """
    Add a subtle background to a table cell.
    """

    tc_pr = cell._tc.get_or_add_tcPr()

    shade = OxmlElement("w:shd")

    shade.set(
        qn("w:fill"),
        color,
    )

    tc_pr.append(shade)


# ---------------------------------------------
# Document layout and typography
# ---------------------------------------------

def _style_document(document):
    """
    Configure page margins, fonts,
    headings, headers and footers.
    """

    section = document.sections[0]

    section.top_margin = Inches(0.92)
    section.bottom_margin = Inches(0.77)

    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)

    section.header_distance = Inches(0.39)
    section.footer_distance = Inches(0.38)

    # -----------------------------------------
    # Normal body text
    # -----------------------------------------

    normal = document.styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = NAVY

    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.17

    # -----------------------------------------
    # Section headings
    # -----------------------------------------

    headings = (
        ("Heading 1", 13.5, 16, 7),
        ("Heading 2", 11.5, 12, 5),
        ("Heading 3", 10.5, 9, 4),
    )

    for name, size, before, after in headings:

        style = document.styles[name]

        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True

        if name == "Heading 1":
            style.font.color.rgb = TEAL
        else:
            style.font.color.rgb = NAVY

        style.paragraph_format.space_before = (
            Pt(before)
        )

        style.paragraph_format.space_after = (
            Pt(after)
        )

        style.paragraph_format.keep_with_next = True

    # -----------------------------------------
    # Page header
    # -----------------------------------------

    header = section.header.paragraphs[0]

    header.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    run = header.add_run(
        "COMPONENT C  /  LEGAL RESEARCH"
    )

    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = SLATE

    _bottom_border(
        header,
        "CBD7DE",
        "5",
    )

    # -----------------------------------------
    # Page footer
    # -----------------------------------------

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    label = footer.add_run(
        "AI-ASSISTED DRAFT   |   PAGE "
    )

    label.font.size = Pt(8)
    label.font.color.rgb = SLATE

    # Dynamic page number
    page_field = OxmlElement(
        "w:fldSimple"
    )

    page_field.set(
        qn("w:instr"),
        "PAGE",
    )

    footer._p.append(page_field)


# ---------------------------------------------
# Inline Markdown formatting
# ---------------------------------------------

def _add_formatted(
    paragraph,
    content,
):
    """
    Convert basic Markdown into Word formatting.

    Supports:
    - Bold text
    - Inline code
    - Highlighted CLAIM IDs
    """

    segments = re.split(
        r"(\*\*.*?\*\*|`[^`]+`|CLAIM_\d+)",
        content,
    )

    for segment in segments:

        if not segment:
            continue

        # Bold text
        if (
            segment.startswith("**")
            and segment.endswith("**")
        ):

            run = paragraph.add_run(
                segment[2:-2]
            )

            run.bold = True

        # Inline code and identifiers
        elif (
            segment.startswith("`")
            and segment.endswith("`")
        ):

            run = paragraph.add_run(
                segment[1:-1]
            )

            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = TEAL

        else:

            run = paragraph.add_run(
                segment
            )

            # Emphasize claim identifiers
            if re.fullmatch(
                r"CLAIM_\d+",
                segment,
            ):

                run.bold = True
                run.font.color.rgb = TEAL


# ---------------------------------------------
# Markdown line conversion
# ---------------------------------------------

def _add_markdown_line(
    document,
    line,
):
    """
    Add a Markdown line to the Word document.
    """

    line = line.strip()

    if not line:
        return

    if line in {"---", "***", "___"}:
        return

    # Third-level heading
    if line.startswith("### "):

        paragraph = document.add_heading(
            level=3
        )

        _add_formatted(
            paragraph,
            line[4:],
        )

    # Second-level heading
    elif line.startswith("## "):

        paragraph = document.add_heading(
            level=2
        )

        _add_formatted(
            paragraph,
            line[3:],
        )

    # First-level heading
    elif line.startswith("# "):

        paragraph = document.add_heading(
            level=1
        )

        _add_formatted(
            paragraph,
            line[2:],
        )

        _bottom_border(
            paragraph,
            "D2DEE3",
            "4",
        )

    # Bullet lists
    elif (
        line.startswith("- ")
        or line.startswith("* ")
    ):

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        _add_formatted(
            paragraph,
            line[2:],
        )

    # Blockquotes
    elif line.startswith("> "):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.left_indent = (
            Inches(0.22)
        )

        _add_formatted(
            paragraph,
            line[2:],
        )

        for run in paragraph.runs:
            run.italic = True

    # Standard text
    else:

        paragraph = document.add_paragraph()

        _add_formatted(
            paragraph,
            line,
        )

        if re.search(
            r"\bCLAIM_\d+\b",
            line,
        ):

            paragraph.paragraph_format.space_before = (
                Pt(4)
            )


# ---------------------------------------------
# Main Word exporter
# ---------------------------------------------

def export_legal_brief_docx(
    brief: LegalBrief,
) -> bytes:
    """
    Generate a professionally formatted DOCX.

    Returns:
        Word document bytes.

    The original LegalBrief is not modified.
    """

    document = Document()

    _style_document(document)

    document.core_properties.title = brief.title

    document.core_properties.subject = (
        "AI-assisted legal research draft"
    )

    # -----------------------------------------
    # Document category
    # -----------------------------------------

    eyebrow = document.add_paragraph()

    eyebrow.paragraph_format.space_after = (
        Pt(5)
    )

    run = eyebrow.add_run(
        "LEGAL RESEARCH BRIEF  /  REVIEW COPY"
    )

    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEAL

    # -----------------------------------------
    # Main document title
    # -----------------------------------------

    title = document.add_paragraph()

    title.paragraph_format.space_after = (
        Pt(13)
    )

    title.paragraph_format.keep_with_next = True

    title_run = title.add_run(
        brief.title
    )

    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = NAVY

    _bottom_border(title)

    # -----------------------------------------
    # Highlighted disclaimer
    # -----------------------------------------

    notice_table = document.add_table(
        rows=1,
        cols=1,
    )

    notice = notice_table.cell(0, 0)

    _shade_cell(notice)

    warning = notice.paragraphs[0]

    warning.paragraph_format.space_before = (
        Pt(6)
    )

    warning.paragraph_format.space_after = (
        Pt(6)
    )

    text_run = warning.add_run(
        "REVIEW REQUIRED  |  " + DISCLAIMER
    )

    text_run.font.size = Pt(9)
    text_run.font.color.rgb = NAVY

    # Spacing before main content
    document.add_paragraph().paragraph_format.space_after = (
        Pt(1)
    )

    # -----------------------------------------
    # Insert the generated legal brief
    # -----------------------------------------

    markdown = format_legal_brief(
        brief
    )

    for raw_line in markdown.splitlines():

        line = raw_line.strip()

        # Avoid repeating the title
        if line in {
            f"# {brief.title}",
            f"## {brief.title}",
        }:
            continue

        _add_markdown_line(
            document,
            raw_line,
        )

    # -----------------------------------------
    # Save Word document to memory
    # -----------------------------------------

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()