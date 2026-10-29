
"""
Component C - DOCX Export Tests

Uses a previously recorded Gemini legal brief.

No live Gemini calls are made.
"""

from io import BytesIO
from pathlib import Path
from zipfile import is_zipfile

from docx import Document

from brief.output_models import LegalBrief
from brief.docx_exporter import (
    export_legal_brief_docx,
)


BRIEF_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "demo"
    / "briefs"
    / "improved_baseline_legal_brief.json"
)


def load_brief():
    return LegalBrief.model_validate_json(
        BRIEF_PATH.read_text(
            encoding="utf-8"
        )
    )


def document_text(document):
    """
    Extract text from the Word document.

    Includes:
    - Normal paragraphs
    - Table cells containing formatted content
    """

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text_parts.append(paragraph.text)

    # Extract text inside tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text_parts.append(paragraph.text)

    return "\n".join(text_parts)


def test_docx_export_produces_valid_word_file():
    brief = load_brief()

    file_bytes = export_legal_brief_docx(
        brief
    )

    assert isinstance(file_bytes, bytes)
    assert len(file_bytes) > 0
    assert is_zipfile(BytesIO(file_bytes))

    document = Document(
        BytesIO(file_bytes)
    )

    assert document is not None


def test_docx_contains_generated_brief_content():
    brief = load_brief()

    file_bytes = export_legal_brief_docx(
        brief
    )

    document = Document(
        BytesIO(file_bytes)
    )

    text = document_text(document)

    assert brief.title in text
    assert brief.legal_issue in text
    assert brief.legal_analysis[0].claim_id in text


def test_docx_contains_disclaimer():
    brief = load_brief()

    document = Document(
        BytesIO(
            export_legal_brief_docx(brief)
        )
    )

    text = document_text(document)

    assert "AI-assisted legal research draft" in text
    assert "independently verified" in text


def test_docx_export_does_not_modify_brief():
    brief = load_brief()

    before = brief.model_dump(
        mode="json"
    )

    export_legal_brief_docx(
        brief
    )

    after = brief.model_dump(
        mode="json"
    )

    assert before == after
