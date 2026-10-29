
"""
Component C - Function 2
Legal Brief Download API Tests

Tests:
- Recorded DOCX download
- Recorded PDF download
- POST DOCX export
- POST PDF export
- Invalid LegalBrief rejection
- Missing recorded brief handling

No Gemini API calls are made.
"""

import json

from io import BytesIO
from pathlib import Path

import pytest

from docx import Document
from pypdf import PdfReader
from fastapi.testclient import TestClient

from api.main import app
import api.brief_api as brief_api


# ---------------------------------------------
# Test configuration
# ---------------------------------------------

BRIEF_FILE = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "demo"
    / "briefs"
    / "improved_baseline_legal_brief.json"
)

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)

PDF_MEDIA_TYPE = "application/pdf"


@pytest.fixture
def client():

    app.dependency_overrides.clear()

    with TestClient(app) as test_client:
        yield test_client

    app.dependency_overrides.clear()


# ---------------------------------------------
# Helper functions
# ---------------------------------------------

def load_brief_json():

    return json.loads(
        BRIEF_FILE.read_text(
            encoding="utf-8"
        )
    )


def get_docx_text(data: bytes) -> str:
    """
    Extract normal paragraphs and table text
    from the generated Word document.
    """

    document = Document(
        BytesIO(data)
    )

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
    ]

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    parts.append(
                        paragraph.text
                    )

    return "\n".join(parts)


def get_pdf_text(data: bytes) -> str:
    """
    Extract text from the generated PDF.
    """

    reader = PdfReader(
        BytesIO(data)
    )

    assert len(reader.pages) >= 1

    return "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )


def validate_download(
    response,
    extension: str,
    media_type: str,
):
    """
    Verify HTTP headers, file bytes,
    and actual document contents.
    """

    assert response.status_code == 200

    assert (
        response.headers["content-type"].split(";")[0]
        == media_type
    )

    assert (
        f'filename="legal_research_brief.{extension}"'
        in response.headers["content-disposition"]
    )

    assert len(response.content) > 0

    if extension == "docx":

        assert response.content.startswith(
            b"PK"
        )

        text = get_docx_text(
            response.content
        )

    else:

        assert response.content.startswith(
            b"%PDF-"
        )

        text = get_pdf_text(
            response.content
        )

    assert "Legal Research Brief" in text

    assert "AI-assisted legal research draft" in text

    assert "independently verified" in text


# ---------------------------------------------
# Test 1-2: Recorded downloads
# ---------------------------------------------

@pytest.mark.parametrize(
    "extension,media_type",
    [
        ("docx", DOCX_MEDIA_TYPE),
        ("pdf", PDF_MEDIA_TYPE),
    ],
)
def test_recorded_document_download(
    client,
    extension,
    media_type,
):

    response = client.get(
        f"/api/brief/demo/download/{extension}"
    )

    validate_download(
        response,
        extension,
        media_type,
    )


# ---------------------------------------------
# Test 3-4: POST document exports
# ---------------------------------------------

@pytest.mark.parametrize(
    "extension,media_type",
    [
        ("docx", DOCX_MEDIA_TYPE),
        ("pdf", PDF_MEDIA_TYPE),
    ],
)
def test_post_document_export(
    client,
    extension,
    media_type,
):

    brief = load_brief_json()

    response = client.post(
        f"/api/brief/export/{extension}",
        json=brief,
    )

    validate_download(
        response,
        extension,
        media_type,
    )


# ---------------------------------------------
# Test 5-6: Invalid input rejection
# ---------------------------------------------

@pytest.mark.parametrize(
    "extension",
    [
        "docx",
        "pdf",
    ],
)
def test_invalid_brief_returns_422(
    client,
    extension,
):

    response = client.post(
        f"/api/brief/export/{extension}",
        json={"invalid": True},
    )

    assert response.status_code == 422


# ---------------------------------------------
# Test 7-8: Missing recorded file
# ---------------------------------------------

@pytest.mark.parametrize(
    "extension",
    [
        "docx",
        "pdf",
    ],
)
def test_missing_recorded_brief_returns_404(
    client,
    monkeypatch,
    tmp_path,
    extension,
):

    missing_file = (
        tmp_path / "missing_brief.json"
    )

    monkeypatch.setattr(
        brief_api,
        "RECORDED_BRIEF_FILE",
        missing_file,
    )

    response = client.get(
        f"/api/brief/demo/download/{extension}"
    )

    assert response.status_code == 404
